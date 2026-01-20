"""
Excel import/export utilities for project service.
"""

import logging
from io import BytesIO
from typing import Any

from django.core.exceptions import ValidationError
from django.db import transaction

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False
    pd = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    Workbook = None

from .models import Client, Project, Task, Milestone, Sprint, Contract

logger = logging.getLogger(__name__)


class ExcelImportExport:
    """Base class for Excel import/export operations"""

    def __init__(self, tenant_id=None):
        self.tenant_id = tenant_id
        self.errors = []
        self.warnings = []

    def add_error(self, message: str, row: int | None = None):
        """Add an error message"""
        if row:
            self.errors.append(f"Row {row}: {message}")
        else:
            self.errors.append(message)

    def add_warning(self, message: str, row: int | None = None):
        """Add a warning message"""
        if row:
            self.warnings.append(f"Row {row}: {message}")
        else:
            self.warnings.append(message)


class ClientExcelHandler(ExcelImportExport):
    """Handle Excel import/export for clients"""

    EXPORT_COLUMNS = [
        "name",
        "email",
        "phone",
        "status",
        "industry",
        "company_size",
        "lead_source",
        "lead_score",
        "created_at",
    ]

    IMPORT_COLUMNS = [
        "name",
        "email",
        "phone",
        "status",
        "industry",
        "company_size",
        "lead_source",
        "lead_score",
    ]

    def export_clients(self) -> BytesIO:
        """Export clients to Excel file"""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export functionality")

        queryset = Client.objects.filter(tenant_id=self.tenant_id) if self.tenant_id else Client.objects.none()
        data = []

        for client in queryset:
            data.append(
                {
                    "name": client.name,
                    "email": client.email,
                    "phone": client.phone or "",
                    "status": client.status,
                    "industry": client.industry or "",
                    "company_size": client.company_size or "",
                    "lead_source": client.lead_source or "",
                    "lead_score": client.lead_score,
                    "created_at": client.created_at.strftime("%Y-%m-%d %H:%M:%S") if client.created_at else "",
                }
            )

        df = pd.DataFrame(data)
        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Clients", index=False)
            worksheet = writer.sheets["Clients"]
            self._format_worksheet(worksheet, len(data), len(self.EXPORT_COLUMNS))

        output.seek(0)
        return output

    def import_clients(self, file) -> dict:
        """Import clients from Excel file"""
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas is required for Excel import functionality")

        try:
            df = pd.read_excel(file)
            created_count = 0
            updated_count = 0
            error_count = 0

            for index, row in df.iterrows():
                row_num = index + 2
                try:
                    data = {col: row.get(col) for col in self.IMPORT_COLUMNS if col in row}

                    if not data.get("name") or not data.get("email"):
                        self.add_error("Name and email are required", row_num)
                        error_count += 1
                        continue

                    if self.tenant_id:
                        data["tenant_id"] = self.tenant_id

                    try:
                        client, created = Client.objects.update_or_create(
                            email=data["email"],
                            defaults=data,
                        )
                        if created:
                            created_count += 1
                        else:
                            updated_count += 1
                    except Exception as e:
                        self.add_error(str(e), row_num)
                        error_count += 1

                except Exception as e:
                    self.add_error(f"Error processing row: {str(e)}", row_num)
                    error_count += 1

            return {
                "created": created_count,
                "updated": updated_count,
                "errors": error_count,
                "warnings": self.warnings,
                "error_details": self.errors,
            }

        except Exception as e:
            logger.error(f"Error importing clients: {e}")
            raise

    def _format_worksheet(self, worksheet, num_rows, num_cols):
        """Format Excel worksheet"""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col in range(1, num_cols + 1):
            col_letter = get_column_letter(col)
            cell = worksheet[f"{col_letter}1"]
            cell.fill = header_fill
            cell.font = header_font
            worksheet.column_dimensions[col_letter].width = 15

        for row in range(1, num_rows + 2):
            for col in range(1, num_cols + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.alignment = Alignment(horizontal="left", vertical="center")


class ProjectExcelHandler(ExcelImportExport):
    """Handle Excel import/export for projects"""

    EXPORT_COLUMNS = [
        "name",
        "status",
        "priority",
        "phase",
        "risk_level",
        "start_date",
        "end_date",
        "budget",
        "progress",
        "created_at",
    ]

    IMPORT_COLUMNS = [
        "name",
        "status",
        "priority",
        "phase",
        "risk_level",
        "start_date",
        "end_date",
        "budget",
    ]

    def export_projects(self) -> BytesIO:
        """Export projects to Excel file"""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export functionality")

        queryset = Project.objects.filter(tenant_id=self.tenant_id) if self.tenant_id else Project.objects.none()
        data = []

        for project in queryset:
            data.append(
                {
                    "name": project.name,
                    "status": project.status,
                    "priority": project.priority,
                    "phase": project.phase,
                    "risk_level": project.risk_level,
                    "start_date": str(project.start_date) if project.start_date else "",
                    "end_date": str(project.end_date) if project.end_date else "",
                    "budget": float(project.budget),
                    "progress": project.progress,
                    "created_at": project.created_at.strftime("%Y-%m-%d %H:%M:%S") if project.created_at else "",
                }
            )

        df = pd.DataFrame(data)
        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Projects", index=False)
            worksheet = writer.sheets["Projects"]
            self._format_worksheet(worksheet, len(data), len(self.EXPORT_COLUMNS))

        output.seek(0)
        return output

    def import_projects(self, file) -> dict:
        """Import projects from Excel file"""
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas is required for Excel import functionality")

        try:
            df = pd.read_excel(file)
            created_count = 0
            updated_count = 0
            error_count = 0

            for index, row in df.iterrows():
                row_num = index + 2
                try:
                    data = {col: row.get(col) for col in self.IMPORT_COLUMNS if col in row}

                    if not data.get("name"):
                        self.add_error("Name is required", row_num)
                        error_count += 1
                        continue

                    if self.tenant_id:
                        data["tenant_id"] = self.tenant_id

                    try:
                        project, created = Project.objects.update_or_create(
                            name=data["name"],
                            tenant_id=data.get("tenant_id"),
                            defaults=data,
                        )
                        if created:
                            created_count += 1
                        else:
                            updated_count += 1
                    except Exception as e:
                        self.add_error(str(e), row_num)
                        error_count += 1

                except Exception as e:
                    self.add_error(f"Error processing row: {str(e)}", row_num)
                    error_count += 1

            return {
                "created": created_count,
                "updated": updated_count,
                "errors": error_count,
                "warnings": self.warnings,
                "error_details": self.errors,
            }

        except Exception as e:
            logger.error(f"Error importing projects: {e}")
            raise

    def _format_worksheet(self, worksheet, num_rows, num_cols):
        """Format Excel worksheet"""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col in range(1, num_cols + 1):
            col_letter = get_column_letter(col)
            cell = worksheet[f"{col_letter}1"]
            cell.fill = header_fill
            cell.font = header_font
            worksheet.column_dimensions[col_letter].width = 15

        for row in range(1, num_rows + 2):
            for col in range(1, num_cols + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.alignment = Alignment(horizontal="left", vertical="center")


class TaskExcelHandler(ExcelImportExport):
    """Handle Excel import/export for tasks"""

    EXPORT_COLUMNS = [
        "title",
        "status",
        "assignee_name",
        "start_date",
        "end_date",
        "estimated_hours",
        "created_at",
    ]

    IMPORT_COLUMNS = [
        "title",
        "status",
        "start_date",
        "end_date",
        "estimated_hours",
    ]

    def export_tasks(self) -> BytesIO:
        """Export tasks to Excel file"""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export functionality")

        queryset = Task.objects.filter(tenant_id=self.tenant_id) if self.tenant_id else Task.objects.none()
        data = []

        for task in queryset:
            data.append(
                {
                    "title": task.title,
                    "status": task.status,
                    "assignee_name": task.assignee_name or "",
                    "start_date": str(task.start_date) if task.start_date else "",
                    "end_date": str(task.end_date) if task.end_date else "",
                    "estimated_hours": task.estimated_hours or 0,
                    "created_at": task.created_at.strftime("%Y-%m-%d %H:%M:%S") if task.created_at else "",
                }
            )

        df = pd.DataFrame(data)
        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Tasks", index=False)
            worksheet = writer.sheets["Tasks"]
            self._format_worksheet(worksheet, len(data), len(self.EXPORT_COLUMNS))

        output.seek(0)
        return output

    def import_tasks(self, file) -> dict:
        """Import tasks from Excel file"""
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas is required for Excel import functionality")

        try:
            df = pd.read_excel(file)
            created_count = 0
            updated_count = 0
            error_count = 0

            for index, row in df.iterrows():
                row_num = index + 2
                try:
                    data = {col: row.get(col) for col in self.IMPORT_COLUMNS if col in row}

                    if not data.get("title"):
                        self.add_error("Title is required", row_num)
                        error_count += 1
                        continue

                    if self.tenant_id:
                        data["tenant_id"] = self.tenant_id

                    try:
                        task, created = Task.objects.update_or_create(
                            title=data["title"],
                            tenant_id=data.get("tenant_id"),
                            defaults=data,
                        )
                        if created:
                            created_count += 1
                        else:
                            updated_count += 1
                    except Exception as e:
                        self.add_error(str(e), row_num)
                        error_count += 1

                except Exception as e:
                    self.add_error(f"Error processing row: {str(e)}", row_num)
                    error_count += 1

            return {
                "created": created_count,
                "updated": updated_count,
                "errors": error_count,
                "warnings": self.warnings,
                "error_details": self.errors,
            }

        except Exception as e:
            logger.error(f"Error importing tasks: {e}")
            raise

    def _format_worksheet(self, worksheet, num_rows, num_cols):
        """Format Excel worksheet"""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col in range(1, num_cols + 1):
            col_letter = get_column_letter(col)
            cell = worksheet[f"{col_letter}1"]
            cell.fill = header_fill
            cell.font = header_font
            worksheet.column_dimensions[col_letter].width = 15

        for row in range(1, num_rows + 2):
            for col in range(1, num_cols + 1):
                cell = worksheet.cell(row=row, column=col)
                cell.alignment = Alignment(horizontal="left", vertical="center")
