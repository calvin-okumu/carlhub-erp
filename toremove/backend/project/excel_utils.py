"""
Excel import/export utilities for DjangoCRM
"""

import logging
from io import BytesIO
from typing import Any

from django.core.exceptions import ValidationError
from django.db import transaction

# Lazy imports for optional dependencies
try:
    import pandas as pd

    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False
    pd = None

try:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    Workbook = None

from .models import Client, Project, Task

logger = logging.getLogger(__name__)


class ExcelImportExport:
    """Base class for Excel import/export operations"""

    def __init__(self, tenant=None):
        self.tenant = tenant
        self.errors = []
        self.warnings = []

    def add_error(self, message: str, row: int | None = None):
        """Add an error message"""
        if row:
            self.errors.append(f"Row {row}: {message}")
        else:
            self.errors.append(message)

    def add_warning(self, message: str, row: int | None = None):
        """Add a warning message"""
        if row:
            self.warnings.append(f"Row {row}: {message}")
        else:
            self.warnings.append(message)


class ClientExcelHandler(ExcelImportExport):
    """Handle Excel import/export for clients"""

    EXPORT_COLUMNS = ["name", "email", "phone", "status", "created_at"]

    IMPORT_COLUMNS = ["name", "email", "phone", "status"]

    def export_clients(self) -> BytesIO:
        """Export clients to Excel file"""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export functionality")

        if self.tenant:
            queryset = Client.objects.filter(tenant=self.tenant)
        else:
            # Security: In development mode with no tenant context, return no data
            queryset = Client.objects.none()

        data = []
        for client in queryset:
            data.append(
                {
                    "name": client.name,
                    "email": client.email,
                    "phone": client.phone or "",
                    "status": client.status,
                    "created_at": (
                        client.created_at.strftime("%Y-%m-%d %H:%M:%S") if client.created_at else ""
                    ),
                }
            )

        df = pd.DataFrame(data)
        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Clients", index=False)

            # Format the worksheet
            worksheet = writer.sheets["Clients"]
            self._format_worksheet(worksheet, len(data))

        output.seek(0)
        return output

    def import_clients(self, file_content: bytes) -> dict[str, Any]:
        """Import clients from Excel file"""
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas is required for Excel import functionality")

        try:
            df = pd.read_excel(BytesIO(file_content))

            # Validate columns
            missing_columns = set(self.IMPORT_COLUMNS) - set(df.columns)
            if missing_columns:
                raise ValidationError(f"Missing required columns: {', '.join(missing_columns)}")

            imported_count = 0
            updated_count = 0

            with transaction.atomic():
                for index, row in df.iterrows():
                    try:
                        # Check if client exists
                        client, created = Client.objects.get_or_create(
                            email=row["email"],
                            defaults={
                                "name": row["name"],
                                "phone": row.get("phone", ""),
                                "status": row.get("status", "prospect"),
                                "tenant": self.tenant,
                            },
                        )

                        if not created:
                            # Update existing client
                            client.name = row["name"]
                            if pd.notna(row.get("phone")):
                                client.phone = row["phone"]
                            if pd.notna(row.get("status")):
                                client.status = row["status"]
                            client.save()
                            updated_count += 1
                        else:
                            imported_count += 1

                    except Exception as e:
                        self.add_error(
                            f"Error importing client {row.get('email', 'unknown')}: {str(e)}",
                            index + 2,
                        )

            return {
                "imported": imported_count,
                "updated": updated_count,
                "errors": self.errors,
                "warnings": self.warnings,
            }

        except Exception as e:
            raise ValidationError(f"Error reading Excel file: {str(e)}") from e

    def _format_worksheet(self, worksheet, data_rows: int):
        """Format the Excel worksheet"""
        # Header styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        for col_num, _column_title in enumerate(self.EXPORT_COLUMNS, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col_num, _column_title in enumerate(self.EXPORT_COLUMNS, 1):
            column_letter = get_column_letter(col_num)
            max_length = len(_column_title)
            for row_num in range(2, data_rows + 2):
                cell_value = worksheet.cell(row=row_num, column=col_num).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)


class ProjectExcelHandler(ExcelImportExport):
    """Handle Excel import/export for projects"""

    EXPORT_COLUMNS = [
        "name",
        "client_name",
        "status",
        "priority",
        "budget",
        "progress",
        "start_date",
        "end_date",
        "created_at",
    ]

    IMPORT_COLUMNS = [
        "name",
        "client_email",
        "status",
        "priority",
        "budget",
        "start_date",
        "end_date",
    ]

    def export_projects(self) -> BytesIO:
        """Export projects to Excel file"""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export functionality")

        if self.tenant:
            queryset = Project.objects.select_related("client").filter(tenant=self.tenant)
        else:
            # Security: In development mode with no tenant context, return no data
            queryset = Project.objects.none()

        data = []
        for project in queryset:
            data.append(
                {
                    "name": project.name,
                    "client_name": project.client.name if project.client else "",
                    "status": project.status,
                    "priority": project.priority,
                    "budget": float(project.budget) if project.budget else 0,
                    "progress": project.progress,
                    "start_date": (
                        project.start_date.strftime("%Y-%m-%d") if project.start_date else ""
                    ),
                    "end_date": project.end_date.strftime("%Y-%m-%d") if project.end_date else "",
                    "created_at": (
                        project.created_at.strftime("%Y-%m-%d %H:%M:%S")
                        if project.created_at
                        else ""
                    ),
                }
            )

        df = pd.DataFrame(data)
        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Projects", index=False)

            # Format the worksheet
            worksheet = writer.sheets["Projects"]
            self._format_worksheet(worksheet, len(data))

        output.seek(0)
        return output

    def import_projects(self, file_content: bytes) -> dict[str, Any]:
        """Import projects from Excel file"""
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas is required for Excel import functionality")

        try:
            df = pd.read_excel(BytesIO(file_content))

            # Validate columns
            missing_columns = set(self.IMPORT_COLUMNS) - set(df.columns)
            if missing_columns:
                raise ValidationError(f"Missing required columns: {', '.join(missing_columns)}")

            imported_count = 0

            with transaction.atomic():
                for index, row in df.iterrows():
                    try:
                        # Find client by email
                        client_email = row.get("client_email")
                        if not client_email or pd.isna(client_email):
                            self.add_error("Client email is required", index + 2)
                            continue

                        try:
                            client = Client.objects.get(email=client_email, tenant=self.tenant)
                        except Client.DoesNotExist:
                            self.add_error(f"Client with email {client_email} not found", index + 2)
                            continue

                        # Create project
                        Project.objects.create(
                            name=row["name"],
                            client=client,
                            status=row.get("status", "planning"),
                            priority=row.get("priority", "medium"),
                            budget=row.get("budget", 0) if pd.notna(row.get("budget")) else 0,
                            start_date=(
                                pd.to_datetime(row.get("start_date")).date()
                                if pd.notna(row.get("start_date"))
                                else None
                            ),
                            end_date=(
                                pd.to_datetime(row.get("end_date")).date()
                                if pd.notna(row.get("end_date"))
                                else None
                            ),
                            tenant=self.tenant,
                        )

                        imported_count += 1

                    except Exception as e:
                        self.add_error(
                            f"Error importing project {row.get('name', 'unknown')}: {str(e)}",
                            index + 2,
                        )

            return {
                "imported": imported_count,
                "updated": 0,
                "errors": self.errors,
                "warnings": self.warnings,
            }

        except Exception as e:
            raise ValidationError(f"Error reading Excel file: {str(e)}") from e

    def _format_worksheet(self, worksheet, data_rows: int):
        """Format the Excel worksheet"""
        # Header styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        for col_num, _column_title in enumerate(self.EXPORT_COLUMNS, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col_num, _column_title in enumerate(self.EXPORT_COLUMNS, 1):
            column_letter = get_column_letter(col_num)
            max_length = len(_column_title)
            for row_num in range(2, data_rows + 2):
                cell_value = worksheet.cell(row=row_num, column=col_num).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)


class TaskExcelHandler(ExcelImportExport):
    """Handle Excel import/export for tasks"""

    EXPORT_COLUMNS = [
        "title",
        "project_name",
        "milestone_name",
        "assignee_email",
        "status",
        "estimated_hours",
        "start_date",
        "end_date",
        "created_at",
    ]

    IMPORT_COLUMNS = [
        "title",
        "project_name",
        "milestone_name",
        "assignee_email",
        "status",
        "estimated_hours",
        "start_date",
        "end_date",
    ]

    def export_tasks(self) -> BytesIO:
        """Export tasks to Excel file"""
        if not PANDAS_AVAILABLE or not OPENPYXL_AVAILABLE:
            raise ImportError("pandas and openpyxl are required for Excel export functionality")

        if self.tenant:
            queryset = Task.objects.select_related("milestone__project", "assignee").filter(
                tenant=self.tenant
            )
        else:
            # Security: In development mode with no tenant context, return no data
            queryset = Task.objects.none()

        data = []
        for task in queryset:
            data.append(
                {
                    "title": task.title,
                    "project_name": (
                        task.milestone.project.name
                        if task.milestone and task.milestone.project
                        else ""
                    ),
                    "milestone_name": task.milestone.name if task.milestone else "",
                    "assignee_email": task.assignee.email if task.assignee else "",
                    "status": task.status,
                    "estimated_hours": task.estimated_hours or "",
                    "start_date": task.start_date.strftime("%Y-%m-%d") if task.start_date else "",
                    "end_date": task.end_date.strftime("%Y-%m-%d") if task.end_date else "",
                    "created_at": (
                        task.created_at.strftime("%Y-%m-%d %H:%M:%S") if task.created_at else ""
                    ),
                }
            )

        df = pd.DataFrame(data)
        output = BytesIO()

        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            df.to_excel(writer, sheet_name="Tasks", index=False)

            # Format the worksheet
            worksheet = writer.sheets["Tasks"]
            self._format_worksheet(worksheet, len(data))

        output.seek(0)
        return output

    def import_tasks(self, file_content: bytes) -> dict[str, Any]:
        """Import tasks from Excel file"""
        if not PANDAS_AVAILABLE:
            raise ImportError("pandas is required for Excel import functionality")

        try:
            from django.contrib.auth import get_user_model

            from .models import Milestone

            User = get_user_model()

            df = pd.read_excel(BytesIO(file_content))

            # Validate columns
            missing_columns = set(self.IMPORT_COLUMNS) - set(df.columns)
            if missing_columns:
                raise ValidationError(f"Missing required columns: {', '.join(missing_columns)}")

            imported_count = 0

            with transaction.atomic():
                for index, row in df.iterrows():
                    try:
                        # Find project by name
                        project_name = row.get("project_name")
                        if not project_name or pd.isna(project_name):
                            self.add_error("Project name is required", index + 2)
                            continue

                        try:
                            project = Project.objects.get(name=project_name, tenant=self.tenant)
                        except Project.DoesNotExist:
                            self.add_error(f"Project '{project_name}' not found", index + 2)
                            continue

                        # Find milestone by name within the project
                        milestone_name = row.get("milestone_name")
                        if not milestone_name or pd.isna(milestone_name):
                            self.add_error("Milestone name is required", index + 2)
                            continue

                        try:
                            milestone = Milestone.objects.get(
                                name=milestone_name, project=project, tenant=self.tenant
                            )
                        except Milestone.DoesNotExist:
                            self.add_error(
                                f"Milestone '{milestone_name}' not found in project '{project_name}'",
                                index + 2,
                            )
                            continue

                        # Find assignee by email (optional)
                        assignee = None
                        assignee_email = row.get("assignee_email")
                        if assignee_email and not pd.isna(assignee_email):
                            try:
                                assignee = User.objects.get(email=assignee_email)
                            except User.DoesNotExist:
                                self.add_warning(
                                    f"Assignee with email {assignee_email} not found, task will be unassigned",
                                    index + 2,
                                )

                        # Create task
                        Task.objects.create(
                            title=row["title"],
                            milestone=milestone,
                            status=row.get("status", "to_do"),
                            assignee=assignee,
                            estimated_hours=(
                                row.get("estimated_hours")
                                if pd.notna(row.get("estimated_hours"))
                                else None
                            ),
                            start_date=(
                                pd.to_datetime(row.get("start_date")).date()
                                if pd.notna(row.get("start_date"))
                                else None
                            ),
                            end_date=(
                                pd.to_datetime(row.get("end_date")).date()
                                if pd.notna(row.get("end_date"))
                                else None
                            ),
                            tenant=self.tenant,
                        )

                        imported_count += 1

                    except Exception as e:
                        self.add_error(
                            f"Error importing task {row.get('title', 'unknown')}: {str(e)}",
                            index + 2,
                        )

            return {
                "imported": imported_count,
                "updated": 0,
                "errors": self.errors,
                "warnings": self.warnings,
            }

        except Exception as e:
            raise ValidationError(f"Error reading Excel file: {str(e)}") from e

    def _format_worksheet(self, worksheet, data_rows: int):
        """Format the Excel worksheet"""
        # Header styling
        header_font = Font(bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        for col_num, _column_title in enumerate(self.EXPORT_COLUMNS, 1):
            cell = worksheet.cell(row=1, column=col_num)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")

        # Auto-adjust column widths
        for col_num, _column_title in enumerate(self.EXPORT_COLUMNS, 1):
            column_letter = get_column_letter(col_num)
            max_length = len(_column_title)
            for row_num in range(2, data_rows + 2):
                cell_value = worksheet.cell(row=row_num, column=col_num).value
                if cell_value:
                    max_length = max(max_length, len(str(cell_value)))
            worksheet.column_dimensions[column_letter].width = min(max_length + 2, 50)
